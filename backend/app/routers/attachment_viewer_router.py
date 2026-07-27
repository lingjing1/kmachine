
import os
import json
import base64
import mammoth
from pathlib import Path
from typing import List, Optional
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse, JSONResponse
from pydantic import BaseModel
from sqlalchemy import text
from pptx import Presentation
from PIL import Image
from backend.app.utils.time_utils import get_now_taipei
import io

from backend.app.utils.db_logger import engine
from backend.app.utils.auth_utils import get_current_user_id
from backend.app.utils.concurrency import run_in_db_pool
from backend.app.config.settings import settings

router = APIRouter(prefix="/api/attachments", tags=["Attachment Viewer"])

# ==================== Utility Functions ====================

def resolve_attachment_path(db_path: str) -> Path:
    """
    從 DB 路徑解析出實際檔案路徑。
    考量到開發環境路徑不同（例如 /home/monica/ vs /home/culture/），
    我們提取 filename 部分，並重新結合伺服器目前的 STORAGE_BASE。
    """
    # 如果路徑中包含標準的 storage 標籤，則提取之後的部分
    storage_tag = "/storage/uploads/"
    idx = db_path.find(storage_tag)
    
    if idx != -1:
        # 取得檔名（含雜湊碼前綴的部分）
        filename = db_path[idx + len(storage_tag):]
        # 使用目前的設定檔路徑進行拼接
        return Path(settings.attachment_dir) / filename
    
    # 備用方案：如果沒有標籤，直接用原始檔名嘗試在 uploads 裡找
    base_filename = os.path.basename(db_path)
    return Path(settings.attachment_dir) / base_filename

class ReadingTimeRequest(BaseModel):
    reading_time_seconds: int

# ==================== Content Conversion Logic ====================

def process_pptx_slides(file_path: Path):
    """
    使用 python-pptx 解析簡報內容
    回傳每一頁的文字與圖片 (Base64)
    """
    prs = Presentation(str(file_path))
    slides_data = []
    
    for i, slide in enumerate(prs.slides):
        slide_info = {
            "slide_index": i + 1,
            "texts": [],
            "images": []
        }
        
        for shape in slide.shapes:
            # 提取文字
            if hasattr(shape, "text"):
                text = shape.text.strip()
                # 過濾掉投影片頁碼佔位符 <#> 以及空字串
                if text and text != '<#>' and not text.startswith('‹#›'):
                    if text not in slide_info["texts"]: # 避免重複抓取
                        slide_info["texts"].append(text)
            
            # 提取圖片
            if shape.shape_type == 13: # 圖片類型
                image = shape.image
                image_bytes = image.blob
                base64_img = base64.b64encode(image_bytes).decode('utf-8')
                slide_info["images"].append({
                    "content": base64_img,
                    "content_type": image.content_type
                })
        
        slides_data.append(slide_info)
    
    return slides_data

# ==================== Endpoints ====================

@router.get("/{attachment_id}/view")
async def view_attachment(attachment_id: int, current_user_id: int = Depends(get_current_user_id)):
    """
    線上預覽附件內容。
    根據檔案類型回傳不同格式。
    """
    def _get_attachment_info(aid):
        with engine.connect() as conn:
            query = text("SELECT file_path, original_file_name, file_type FROM attachments WHERE id = :id")
            return conn.execute(query, {"id": aid}).fetchone()

    row = await run_in_db_pool(_get_attachment_info, attachment_id)
    if not row:
        raise HTTPException(status_code=404, detail="附件不存在")
    
    file_path = resolve_attachment_path(row.file_path)
    if not file_path.exists():
         # Retry with filename in upload_dir as backup
         file_path = Path(settings.attachment_dir) / row.original_file_name

    if not file_path.exists():
        raise HTTPException(status_code=404, detail="檔案實體不存在")

    ext = file_path.suffix.lower()
    
    # 1. PDF 處理 (Streaming)
    if ext == '.pdf':
        def iterfile():
            with open(file_path, mode="rb") as f:
                yield from f
        return StreamingResponse(iterfile(), media_type="application/pdf")

    # 2. Word 處理 (mammoth -> HTML)
    elif ext in ['.doc', '.docx']:
        try:
            with open(file_path, "rb") as docx_file:
                result = mammoth.convert_to_html(docx_file)
                html = result.value
                return JSONResponse(content={"type": "html", "content": html})
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Word 轉檔失敗: {str(e)}")

    # 3. PPTX 處理 (python-pptx -> JSON Slides)
    elif ext == '.pptx':
        try:
            slides = await run_in_threadpool_wrapped(process_pptx_slides, file_path)
            return JSONResponse(content={"type": "pptx", "content": slides})
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"PPTX 轉檔失敗: {str(e)}")

    # 4. 純文字處理 (直接讀取)
    elif ext in ['.txt', '.md', '.csv', '.log']:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                text_content = f.read()
            return JSONResponse(content={"type": "text", "content": text_content})
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"文字檔讀取失敗: {str(e)}")

    # 5. 圖片處理 (Base64)
    elif ext in ['.jpg', '.jpeg', '.png', '.gif', '.webp']:
        try:
            with open(file_path, "rb") as f:
                image_bytes = f.read()
            base64_img = base64.b64encode(image_bytes).decode('utf-8')
            mime_map = {
                '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg',
                '.png': 'image/png', '.gif': 'image/gif', '.webp': 'image/webp'
            }
            mime = mime_map.get(ext, 'image/jpeg')
            return JSONResponse(content={
                "type": "image",
                "content": base64_img,
                "mime": mime
            })
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"圖片讀取失敗: {str(e)}")

    # 6. 不支援線上預覽的格式 → 返回下載提示，讓前端顯示下載按鈕
    else:
        return JSONResponse(content={
            "type": "download",
            "content": row.original_file_name,
            "download_url": f"/api/attachments/{attachment_id}/download"
        })


async def run_in_threadpool_wrapped(func, *args, **kwargs):
    from starlette.concurrency import run_in_threadpool
    return await run_in_threadpool(func, *args, **kwargs)

@router.post("/{attachment_id}/reading-time")
async def log_reading_time(
    attachment_id: int, 
    req: ReadingTimeRequest,
    current_user_id: int = Depends(get_current_user_id)
):
    """
    記錄學生對附件的閱讀時間
    """
    def _sync_log_time(aid, uid, seconds):
        with engine.connect() as conn:
            
            # 使用 ON CONFLICT 更新時間
            upsert_query = text("""
                INSERT INTO attachment_reading_logs (attachment_id, student_id, reading_time_seconds, last_read_at)
                VALUES (:aid, :uid, :sec, :now)
                ON CONFLICT (attachment_id, student_id) 
                DO UPDATE SET 
                    reading_time_seconds = attachment_reading_logs.reading_time_seconds + EXCLUDED.reading_time_seconds,
                    last_read_at = EXCLUDED.last_read_at
            """)
            conn.execute(upsert_query, {"aid": aid, "uid": uid, "sec": seconds, "now": get_now_taipei()})
            conn.commit()
            return True

    await run_in_db_pool(_sync_log_time, attachment_id, current_user_id, req.reading_time_seconds)
    return {"message": "Success"}
